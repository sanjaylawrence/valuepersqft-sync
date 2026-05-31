# =================================================================
# cleaners/call_data.py — Call Data Cleaner
# =================================================================
# Reads ALL Excel files from Google Drive CALL_REPORT folder
# Uses date-range based mobile number mapping
# Handles number reassignment between employees
# Primary key = report_date + mobile_number
# =================================================================

import pandas as pd
import numpy as np
import re
import io
from openpyxl import load_workbook
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload


# ─────────────────────────────────────────────────────────────────
# CONSTANTS
# ─────────────────────────────────────────────────────────────────

PARENT_FOLDER_NAME   = "Valuepersqft_DB"
CALL_FOLDER_NAME     = "CALL_REPORT"
MAPPING_SHEET_NAME   = "MOBILE_NO_MAPPING"
MAPPING_WORKSHEET    = "MOBILE_NO_MAPPING"


# ─────────────────────────────────────────────────────────────────
# GOOGLE DRIVE HELPERS
# ─────────────────────────────────────────────────────────────────

def get_folder_id(drive_service, folder_name, parent_id=None):
    query = f"name='{folder_name}' and mimeType='application/vnd.google-apps.folder' and trashed=false"
    if parent_id:
        query += f" and '{parent_id}' in parents"
    result = drive_service.files().list(q=query, fields="files(id, name)").execute()
    files  = result.get("files", [])
    if not files:
        raise Exception(f"Folder '{folder_name}' not found in Google Drive")
    return files[0]["id"]


def get_excel_files_in_folder(drive_service, folder_id):
    query  = (
        f"'{folder_id}' in parents and trashed=false and "
        f"mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'"
    )
    result = drive_service.files().list(
        q=query,
        fields="files(id, name, modifiedTime)",
        orderBy="modifiedTime desc"
    ).execute()
    return result.get("files", [])


def download_excel(drive_service, file_id):
    request    = drive_service.files().get_media(fileId=file_id)
    buffer     = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request)
    done       = False
    while not done:
        _, done = downloader.next_chunk()
    buffer.seek(0)
    return buffer


# ─────────────────────────────────────────────────────────────────
# MOBILE NUMBER MAPPING — WITH DATE RANGE
# ─────────────────────────────────────────────────────────────────

def get_mapping(sheets_client):
    """
    Read MOBILE_NO_MAPPING sheet with date range columns.
    Returns DataFrame with:
    Mobile_Number | Employee_Name | Start_Date | End_Date
    """
    try:
        sh   = sheets_client.open(MAPPING_SHEET_NAME)
        ws   = sh.worksheet(MAPPING_WORKSHEET)
        data = ws.get_all_records()
        df   = pd.DataFrame(data)
        df.columns = df.columns.str.strip()

        # Find column names dynamically
        phone_col = [c for c in df.columns if 'mobile' in c.lower() or 'phone' in c.lower()][0]
        name_col  = [c for c in df.columns if 'employee' in c.lower() or 'name' in c.lower()][0]
        start_col = [c for c in df.columns if 'start' in c.lower()][0]
        end_col   = [c for c in df.columns if 'end' in c.lower()][0]

        df = df.rename(columns={
            phone_col : 'Mobile_Number',
            name_col  : 'Employee_Name',
            start_col : 'Start_Date',
            end_col   : 'End_Date',
        })

        # Clean mobile numbers — extract 10 digits
        df['Mobile_Number'] = df['Mobile_Number'].astype(str).str.extract(r'(\d{10})')[0]
        df['Employee_Name'] = df['Employee_Name'].astype(str).str.strip().str.title()

        # Parse dates
        df['Start_Date'] = pd.to_datetime(df['Start_Date'], errors='coerce', dayfirst=True)
        df['End_Date']   = pd.to_datetime(df['End_Date'],   errors='coerce', dayfirst=True)

        df = df.dropna(subset=['Mobile_Number'])
        print(f"    ✅ Loaded {len(df)} mobile mappings")
        return df

    except Exception as e:
        print(f"    ⚠️  Could not load mobile mapping: {e}")
        return pd.DataFrame(columns=['Mobile_Number', 'Employee_Name', 'Start_Date', 'End_Date'])


def get_employee_name(mobile, report_date, mapping_df):
    """
    Find correct employee name for a mobile number on a specific date.
    Uses date range: Start_Date <= report_date <= End_Date (or blank End_Date = still active)
    """
    if pd.isna(mobile) or mapping_df.empty:
        return None

    # Filter rows matching this mobile number
    matches = mapping_df[mapping_df['Mobile_Number'] == str(mobile).strip()]

    if matches.empty:
        return None

    report_dt = pd.to_datetime(report_date)

    for _, row in matches.iterrows():
        start = row['Start_Date']
        end   = row['End_Date']

        # Check if report_date falls in range
        start_ok = pd.isna(start) or report_dt >= start
        end_ok   = pd.isna(end)   or report_dt <= end

        if start_ok and end_ok:
            return row['Employee_Name']

    return None


# ─────────────────────────────────────────────────────────────────
# REPORT DATE EXTRACTION
# ─────────────────────────────────────────────────────────────────

def get_report_date(buffer):
    """Extract report date from Excel file header rows."""
    buffer.seek(0)
    wb = load_workbook(buffer, data_only=True)
    ws = wb["Employee Summary"]

    for row in ws.iter_rows(1, 3, values_only=True):
        for cell in row:
            match = re.search(r"\d{1,2}\s\w+\s\d{4}", str(cell))
            if match:
                return pd.to_datetime(match.group(), format="%d %b %Y").date()

    raise ValueError("Date not found in file header")


# ─────────────────────────────────────────────────────────────────
# DURATION CONVERSION
# ─────────────────────────────────────────────────────────────────

def duration_to_mmss(td):
    """Convert timedelta to MM.SS float format."""
    if pd.isna(td):
        return 0.00
    total_seconds = int(td.total_seconds())
    minutes       = total_seconds // 60
    seconds       = total_seconds % 60
    return float(f"{minutes}.{seconds:02d}")


# ─────────────────────────────────────────────────────────────────
# SAFE VALUE HELPERS
# ─────────────────────────────────────────────────────────────────

def safe_str(val):
    if val is None:
        return None
    if isinstance(val, float) and np.isnan(val):
        return None
    s = str(val).strip()
    if s in ("", "nan", "None", "NaT", "NaN"):
        return None
    return s


def safe_num(val):
    if val is None:
        return None
    try:
        f = float(val)
        return None if np.isnan(f) else f
    except (ValueError, TypeError):
        return None


def remove_all_nan(df: pd.DataFrame) -> pd.DataFrame:
    df = df.replace({np.nan: None, float('inf'): None, float('-inf'): None})
    df = df.where(pd.notnull(df), None)
    return df


# ─────────────────────────────────────────────────────────────────
# MAIN CLEAN FUNCTION
# ─────────────────────────────────────────────────────────────────

def clean(df: pd.DataFrame, creds, sheets_client, supabase=None) -> pd.DataFrame:

    print("    🧹 Starting call data cleaning...")

    # ── Build Drive service ───────────────────────────────────
    drive_service = build("drive", "v3", credentials=creds)

    # ── Get folder ────────────────────────────────────────────
    print("    📂 Finding Google Drive folder...")
    parent_id = get_folder_id(drive_service, PARENT_FOLDER_NAME)
    folder_id = get_folder_id(drive_service, CALL_FOLDER_NAME, parent_id=parent_id)

    # ── Get Excel files ───────────────────────────────────────
    excel_files = get_excel_files_in_folder(drive_service, folder_id)
    if not excel_files:
        raise Exception(f"No Excel files found in '{CALL_FOLDER_NAME}'")
    print(f"    📊 Found {len(excel_files)} Excel files")

    # ── Load mobile mapping ───────────────────────────────────
    print("    📋 Loading mobile number mapping...")
    mapping_df = get_mapping(sheets_client)

    # ── Process each Excel file ───────────────────────────────
    all_data = []

    for file_info in excel_files:
        file_name = file_info["name"]
        file_id   = file_info["id"]
        print(f"    📄 Processing: {file_name}")

        try:
            # Download file
            buffer = download_excel(drive_service, file_id)

            # Get report date from header
            report_date = get_report_date(buffer)
            buffer.seek(0)

            # Read Employee Summary sheet with merged headers
            df_file = pd.read_excel(
                buffer,
                sheet_name="Employee Summary",
                header=[2, 3]
            )

            # Flatten multi-level column names
            df_file.columns = [
                f"{a}_{b}".strip("_").replace(" ", "_")
                for a, b in df_file.columns
            ]

            # Find employee and connected calls columns dynamically
            emp_col = next(
                (c for c in df_file.columns if 'employee' in c.lower()), None
            )
            conn_col = next(
                (c for c in df_file.columns
                 if 'connected' in c.lower()
                 and 'call' in c.lower()
                 and 'incoming' not in c.lower()
                 and 'outgoing' not in c.lower()), None
            )

            if not emp_col or not conn_col:
                print(f"    ⚠️  Could not find required columns in {file_name}")
                continue

            # Extract mobile number from employee column
            df_file['Mobile_Number'] = (
                df_file[emp_col].astype(str).str.extract(r'(\d{10})')[0]
            )

            # Map employee name using date range
            df_file['Employee_Name'] = df_file.apply(
                lambda row: get_employee_name(
                    row['Mobile_Number'], report_date, mapping_df
                ),
                axis=1
            )

            # Add report date
            df_file['Report_Date'] = report_date

            # Convert total duration
            df_file['Total_Duration'] = (
                pd.to_timedelta(df_file['Total_Duration'], errors='coerce')
                .apply(duration_to_mmss)
            )

            # Keep only required columns
            df_clean = df_file[[
                'Report_Date', 'Employee_Name', 'Mobile_Number',
                'Total_Call', conn_col, 'Total_Duration'
            ]].rename(columns={conn_col: 'Connected_Calls'})

            # Drop rows with no mobile number
            df_clean = df_clean[df_clean['Mobile_Number'].notna()]

            all_data.append(df_clean)
            print(f"    ✅ {len(df_clean)} rows from {file_name} — Date: {report_date}")

        except Exception as e:
            print(f"    ❌ Error in {file_name}: {e}")
            import traceback
            traceback.print_exc()
            continue

    if not all_data:
        raise Exception("No data processed from any file")

    # ── Combine all files ─────────────────────────────────────
    print("    🔗 Combining all files...")
    final_df = pd.concat(all_data, ignore_index=True)
    print(f"    📊 Total rows: {len(final_df)}")

    # ── Create primary key ────────────────────────────────────
    final_df['call_id'] = (
        final_df['Report_Date'].astype(str).str.strip()
        + "_"
        + final_df['Mobile_Number'].astype(str).str.strip()
    )

    # ── Rename to Supabase column names ──────────────────────
    final_df = final_df.rename(columns={
        'Report_Date'     : 'report_date',
        'Employee_Name'   : 'employee_name',
        'Mobile_Number'   : 'mobile_number',
        'Total_Call'      : 'total_call',
        'Connected_Calls' : 'connected_calls',
        'Total_Duration'  : 'total_duration',
    })

    # ── Keep only Supabase columns ────────────────────────────
    final_columns = [
        'call_id', 'report_date', 'employee_name',
        'mobile_number', 'total_call', 'connected_calls', 'total_duration'
    ]
    final_df = final_df[[c for c in final_columns if c in final_df.columns]]

    # ── Remove duplicates — keep last ─────────────────────────
    before_dedup = len(final_df)
    final_df     = final_df.drop_duplicates(subset=['call_id'], keep='last')
    after_dedup  = len(final_df)
    if before_dedup != after_dedup:
        print(f"    ⚠️  Removed {before_dedup - after_dedup} duplicate call_ids")

    # ── Safe type conversion ──────────────────────────────────
    str_cols = ['call_id', 'employee_name', 'mobile_number']
    num_cols = ['total_call', 'connected_calls', 'total_duration']

    for col in str_cols:
        if col in final_df.columns:
            final_df[col] = final_df[col].apply(safe_str)

    for col in num_cols:
        if col in final_df.columns:
            final_df[col] = final_df[col].apply(safe_num)

    # report_date to string
    if 'report_date' in final_df.columns:
        final_df['report_date'] = pd.to_datetime(
            final_df['report_date'], errors='coerce'
        ).dt.strftime('%Y-%m-%d')
        final_df['report_date'] = final_df['report_date'].where(
            final_df['report_date'].notna(), None
        )

    # ── Final NaN removal ─────────────────────────────────────
    final_df = remove_all_nan(final_df)

    # ── Drop rows with no call_id ─────────────────────────────
    final_df = final_df[final_df['call_id'].notna()]

    print(f"    ✅ Call data cleaning complete — {len(final_df)} rows ready")
    return final_df